from itertools import chain
from operator import attrgetter
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.core.paginator import Paginator
from django.db.models import DecimalField, ExpressionWrapper, F, Q, Sum
from django.db import transaction
from django.db.models.deletion import ProtectedError
from django.http import JsonResponse
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme

from .forms import (
    EmployeeForm, ExpenseCategoryForm, ExpenseForm, OtherIncomeForm, PartCatalogForm,
    RepairFinanceForm, SalaryPaymentForm, StockReceiptForm, SupplierForm,
    SupplierPaymentForm, WarrantyClaimForm, SupplierReturnForm,
    DistributedExpenseForm, PayrollPeriodForm, MasterAccessCreateForm, MasterPasswordForm,
)
from .models import (
    Employee, Expense, ExpenseCategory, OtherIncome, PartCatalog, PartItem,
    RepairFinance, SalaryPayment, StockReceipt, Supplier, SupplierPayment, SupplierReturn, WarrantyClaim,
    DistributedExpense, PayrollPeriod,
)
from .services import close_payroll_period, employee_rows, financial_summary, monthly_chart, payroll_preview, resolve_period, stock_summary, supplier_summary
from .supplier_ledger import supplier_ledger


finance_staff_required = user_passes_test(
    lambda user: user.is_active and user.is_staff and user.has_module_perms("finance"),
    login_url="admin:login",
)


def _page(queryset, request, size=40):
    return Paginator(queryset, size).get_page(request.GET.get("page"))


def _period_context(request):
    start, end, preset = resolve_period(request)
    return start, end, {"date_from": start, "date_to": end, "period": preset}


def _form_view(request, form_class, title, success_message, return_name, instance=None, show_calculator=False, initial=None):
    requested_next = request.POST.get("next") or request.GET.get("next")
    return_url = requested_next if requested_next and url_has_allowed_host_and_scheme(requested_next, allowed_hosts={request.get_host()}) else None
    form = form_class(request.POST or None, instance=instance, initial=initial)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, success_message)
        return redirect(return_url or return_name)
    return render(request, "finance/form.html", {
        "form": form, "title": title, "return_name": return_name, "return_url": return_url, "show_calculator": show_calculator,
    })


def _delete_view(request, obj, title, return_name):
    if request.method == "POST":
        try:
            obj.delete()
            messages.success(request, f"{title} удалено.")
        except ProtectedError:
            messages.error(request, "Запись используется в операциях. Сделайте её неактивной вместо удаления.")
        return redirect(return_name)
    return render(request, "finance/confirm_delete.html", {
        "object": obj, "title": f"Удалить: {title}", "return_name": return_name,
    })


@finance_staff_required
def dashboard(request):
    start, end, period_ctx = _period_context(request)
    summary = financial_summary(start, end)
    expenses_by_category = list(
        Expense.objects.filter(date__range=(start, end)).values("category__name")
        .annotate(total=Sum("amount")).order_by("category__sort_order", "category__name")
    )
    for row in expenses_by_category:
        row["percent"] = round((row["total"] / summary["expenses"] * 100), 1) if summary["expenses"] else 0
    operations = sorted(chain(
        RepairFinance.objects.filter(date__range=(start, end))[:8],
        Expense.objects.filter(date__range=(start, end))[:8],
        OtherIncome.objects.filter(date__range=(start, end))[:8],
    ), key=attrgetter("date", "created_at"), reverse=True)[:12]
    for operation in operations:
        if isinstance(operation, RepairFinance):
            operation.kind, operation.display_amount = "Ремонт", operation.workshop_profit
        elif isinstance(operation, Expense):
            operation.kind, operation.display_amount = "Расход", -operation.amount
        else:
            operation.kind, operation.display_amount = "Доход", operation.amount
    chart = monthly_chart(end.year)
    context = {
        **period_ctx, "summary": summary, "expenses_by_category": expenses_by_category,
        "stock_summary": stock_summary(start, end),
        "employee_rows": employee_rows(start, end), "operations": operations,
        "chart_labels": [row["month"] for row in chart],
        "chart_revenue": [float(row["revenue"]) for row in chart],
        "chart_costs": [float(row["parts"] + row["salary"] + row["expenses"]) for row in chart],
        "chart_profit": [float(row["net_profit"]) for row in chart],
        "distributed_active": DistributedExpense.objects.filter(is_active=True).count(),
        "distributed_period_total": sum((item.per_period_amount for item in DistributedExpense.objects.filter(is_active=True)), Decimal("0.00")),
    }
    return render(request, "finance/dashboard.html", context)


@finance_staff_required
def repair_list(request):
    start, end, context = _period_context(request)
    qs = RepairFinance.objects.select_related("employee").filter(date__range=(start, end))
    employee_id = request.GET.get("employee")
    query = (request.GET.get("q") or "").strip()
    sort = request.GET.get("sort", "-date")
    allowed = {"date", "-date", "revenue", "-revenue", "workshop_profit", "-workshop_profit", "description", "-description"}
    if employee_id:
        qs = qs.filter(employee_id=employee_id)
    if query:
        qs = qs.filter(Q(description__icontains=query) | Q(comment__icontains=query))
    context.update(page_obj=_page(qs.order_by(sort if sort in allowed else "-date", "-created_at"), request), employees=Employee.objects.all(), selected_employee=employee_id, q=query, sort=sort)
    return render(request, "finance/repair_list.html", context)


@finance_staff_required
def repair_create(request):
    return _form_view(request, RepairFinanceForm, "Добавить ремонт", "Ремонт сохранён.", "finance:repair_list", show_calculator=True)


@finance_staff_required
def repair_edit(request, pk):
    return _form_view(request, RepairFinanceForm, "Изменить ремонт", "Ремонт обновлён.", "finance:repair_list", get_object_or_404(RepairFinance, pk=pk), True)


@finance_staff_required
def repair_delete(request, pk):
    return _delete_view(request, get_object_or_404(RepairFinance, pk=pk), "Ремонт", "finance:repair_list")


@finance_staff_required
def expense_list(request):
    start, end, context = _period_context(request)
    qs = Expense.objects.select_related("category").filter(date__range=(start, end))
    category_id, query = request.GET.get("category"), (request.GET.get("q") or "").strip()
    if category_id:
        qs = qs.filter(category_id=category_id)
    if query:
        qs = qs.filter(Q(description__icontains=query) | Q(comment__icontains=query))
    context.update(page_obj=_page(qs, request), categories=ExpenseCategory.objects.all(), selected_category=category_id, q=query)
    return render(request, "finance/expense_list.html", context)


@finance_staff_required
def expense_create(request):
    return _form_view(request, ExpenseForm, "Добавить расход", "Расход сохранён.", "finance:expense_list")


@finance_staff_required
def expense_edit(request, pk):
    return _form_view(request, ExpenseForm, "Изменить расход", "Расход обновлён.", "finance:expense_list", get_object_or_404(Expense, pk=pk))


@finance_staff_required
def expense_delete(request, pk):
    return _delete_view(request, get_object_or_404(Expense, pk=pk), "Расход", "finance:expense_list")


@finance_staff_required
def category_list(request):
    return render(request, "finance/category_list.html", {"categories": ExpenseCategory.objects.all()})


@finance_staff_required
def category_create(request):
    return _form_view(request, ExpenseCategoryForm, "Добавить категорию", "Категория сохранена.", "finance:category_list")


@finance_staff_required
def category_edit(request, pk):
    return _form_view(request, ExpenseCategoryForm, "Изменить категорию", "Категория обновлена.", "finance:category_list", get_object_or_404(ExpenseCategory, pk=pk))


@finance_staff_required
def category_delete(request, pk):
    return _delete_view(request, get_object_or_404(ExpenseCategory, pk=pk), "Категория", "finance:category_list")


@finance_staff_required
def income_list(request):
    start, end, context = _period_context(request)
    query = (request.GET.get("q") or "").strip()
    qs = OtherIncome.objects.filter(date__range=(start, end))
    if query:
        qs = qs.filter(Q(description__icontains=query) | Q(comment__icontains=query))
    context.update(page_obj=_page(qs, request), q=query)
    return render(request, "finance/income_list.html", context)


@finance_staff_required
def income_create(request):
    return _form_view(request, OtherIncomeForm, "Добавить доход", "Доход сохранён.", "finance:income_list")


@finance_staff_required
def income_edit(request, pk):
    return _form_view(request, OtherIncomeForm, "Изменить доход", "Доход обновлён.", "finance:income_list", get_object_or_404(OtherIncome, pk=pk))


@finance_staff_required
def income_delete(request, pk):
    return _delete_view(request, get_object_or_404(OtherIncome, pk=pk), "Доход", "finance:income_list")


@finance_staff_required
def master_list(request):
    start, end, context = _period_context(request)
    context.update(rows=employee_rows(start, end))
    return render(request, "finance/master_list.html", context)


@finance_staff_required
@transaction.atomic
def master_access_create(request, pk):
    employee = get_object_or_404(Employee.objects.select_for_update(), pk=pk)
    if employee.user_id:
        messages.error(request, "У этого мастера уже есть аккаунт.")
        return redirect("finance:master_list")
    form = MasterAccessCreateForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        employee.user = form.save()
        employee.save(update_fields=["user"])
        messages.success(request, f"Доступ для {employee.name} создан.")
        return redirect("finance:master_list")
    return render(request, "finance/master_access_form.html", {"form": form, "employee": employee, "title": "Создать доступ мастера"})


@finance_staff_required
def master_access_password(request, pk):
    employee = get_object_or_404(Employee.objects.select_related("user"), pk=pk, user__isnull=False)
    form = MasterPasswordForm(employee.user, request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, f"Пароль для {employee.name} изменён.")
        return redirect("finance:master_list")
    return render(request, "finance/master_access_form.html", {"form": form, "employee": employee, "title": "Сменить пароль"})


@finance_staff_required
def master_access_toggle(request, pk):
    employee = get_object_or_404(Employee.objects.select_related("user"), pk=pk, user__isnull=False)
    if request.method == "POST":
        employee.user.is_active = not employee.user.is_active
        employee.user.save(update_fields=["is_active"])
        messages.success(request, "Доступ включён." if employee.user.is_active else "Доступ отключён.")
    return redirect("finance:master_list")


@finance_staff_required
def master_create(request):
    return _form_view(request, EmployeeForm, "Добавить мастера", "Мастер сохранён.", "finance:master_list")


@finance_staff_required
def master_edit(request, pk):
    return _form_view(request, EmployeeForm, "Изменить мастера", "Мастер обновлён.", "finance:master_list", get_object_or_404(Employee, pk=pk))


@finance_staff_required
def master_delete(request, pk):
    return _delete_view(request, get_object_or_404(Employee, pk=pk), "Мастер", "finance:master_list")


@finance_staff_required
def salary_list(request):
    start, end, context = _period_context(request)
    employee_id = request.GET.get("employee")
    employee = get_object_or_404(Employee, pk=employee_id) if employee_id else None
    payments = SalaryPayment.objects.select_related("employee").filter(date__range=(start, end))
    if employee:
        payments = payments.filter(employee=employee)
    open_period = PayrollPeriod.objects.filter(status=PayrollPeriod.Status.OPEN).order_by("-start_date").first()
    current_payroll = []
    if open_period:
        current_payroll = [{"employee": person, **payroll_preview(open_period, person)} for person in Employee.objects.filter(is_active=True, is_owner=False)]
    context.update(rows=employee_rows(start, end, employee), page_obj=_page(payments, request), employees=Employee.objects.filter(is_owner=False), selected_employee=employee, open_period=open_period, current_payroll=current_payroll)
    return render(request, "finance/salary_list.html", context)


@finance_staff_required
def salary_payment_create(request):
    initial = {"employee": request.GET.get("employee")} if request.GET.get("employee") else None
    return _form_view(request, SalaryPaymentForm, "Выплата мастеру", "Выплата сохранена.", "finance:salary_list", initial=initial)


@finance_staff_required
def salary_payment_edit(request, pk):
    return _form_view(request, SalaryPaymentForm, "Изменить выплату", "Выплата обновлена.", "finance:salary_list", get_object_or_404(SalaryPayment, pk=pk))


@finance_staff_required
def salary_payment_delete(request, pk):
    return _delete_view(request, get_object_or_404(SalaryPayment, pk=pk), "Выплата", "finance:salary_list")


@finance_staff_required
def employee_percent(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    return JsonResponse({"percent": "0.00" if employee.is_owner else f"{employee.default_percent:.2f}", "is_owner": employee.is_owner})


@finance_staff_required
def stock_list(request):
    query = (request.GET.get("q") or "").strip()
    status = request.GET.get("status", PartItem.Status.IN_STOCK)
    qs = PartItem.objects.select_related("receipt__part", "receipt__supplier").prefetch_related("warranty_claims")
    if status in PartItem.Status.values:
        qs = qs.filter(status=status)
    if query:
        qs = qs.filter(
            Q(receipt__part__brand__icontains=query) | Q(receipt__part__device_model__icontains=query)
            | Q(receipt__part__name__icontains=query) | Q(receipt__supplier__name__icontains=query)
            | Q(receipt__order_number__icontains=query)
        )
        if query.upper().lstrip("#P").isdigit():
            qs = qs | PartItem.objects.filter(pk=int(query.upper().lstrip("#P"))).select_related("receipt__part", "receipt__supplier")
    return render(request, "finance/stock_list.html", {"page_obj": _page(qs.distinct(), request), "q": query, "status": status, "statuses": PartItem.Status.choices})


@finance_staff_required
def part_item_detail(request, pk):
    item = get_object_or_404(PartItem.objects.select_related("receipt__part", "receipt__supplier").prefetch_related("warranty_claims"), pk=pk)
    usage = getattr(item, "repair_usage", None)
    return render(request, "finance/part_item_detail.html", {"item": item, "usage": usage})


@finance_staff_required
def receipt_list(request):
    start, end, context = _period_context(request)
    qs = StockReceipt.objects.select_related("supplier", "part").filter(date__range=(start, end))
    supplier_id, query = request.GET.get("supplier"), (request.GET.get("q") or "").strip()
    if supplier_id:
        qs = qs.filter(supplier_id=supplier_id)
    if query:
        qs = qs.filter(Q(part__brand__icontains=query) | Q(part__device_model__icontains=query) | Q(part__name__icontains=query) | Q(order_number__icontains=query))
    total_value = sum((row.total_cost for row in qs), 0)
    context.update(page_obj=_page(qs, request), suppliers=Supplier.objects.all(), selected_supplier=supplier_id, q=query, total_value=total_value)
    return render(request, "finance/receipt_list.html", context)


@finance_staff_required
def receipt_create(request):
    return _form_view(request, StockReceiptForm, "Добавить поступление", "Поступление сохранено, экземпляры созданы.", "finance:receipt_list")


@finance_staff_required
def receipt_edit(request, pk):
    return _form_view(request, StockReceiptForm, "Изменить поступление", "Поступление обновлено.", "finance:receipt_list", get_object_or_404(StockReceipt, pk=pk))


@finance_staff_required
def catalog_list(request):
    query = (request.GET.get("q") or "").strip()
    qs = PartCatalog.objects.all()
    if query:
        qs = qs.filter(Q(brand__icontains=query) | Q(device_model__icontains=query) | Q(name__icontains=query))
    return render(request, "finance/catalog_list.html", {"page_obj": _page(qs, request), "q": query})


@finance_staff_required
def catalog_create(request):
    return _form_view(request, PartCatalogForm, "Добавить наименование", "Запчасть добавлена в каталог.", "finance:catalog_list")


@finance_staff_required
def catalog_edit(request, pk):
    return _form_view(request, PartCatalogForm, "Изменить наименование", "Каталог обновлён.", "finance:catalog_list", get_object_or_404(PartCatalog, pk=pk))


@finance_staff_required
def supplier_list(request):
    query = (request.GET.get("q") or "").strip()
    qs = Supplier.objects.all()
    if query:
        qs = qs.filter(Q(name__icontains=query) | Q(contact_person__icontains=query) | Q(phone__icontains=query))
    suppliers = list(qs)
    total_debt = Decimal("0.00")
    total_credit = Decimal("0.00")
    for supplier in suppliers:
        ledger_summary = supplier_ledger(supplier)["summary"]
        supplier.received_total = ledger_summary["received"]
        supplier.paid_total = ledger_summary["paid"]
        supplier.account_balance = -ledger_summary["current"]
        supplier.debt_amount = max(-supplier.account_balance, Decimal("0.00"))
        supplier.credit_amount = max(supplier.account_balance, Decimal("0.00"))
        total_debt += supplier.debt_amount
        total_credit += supplier.credit_amount
    return render(request, "finance/supplier_list.html", {
        "suppliers": suppliers, "q": query,
        "total_debt": total_debt, "total_credit": total_credit,
        "net_balance": total_credit - total_debt,
    })


@finance_staff_required
def supplier_create(request):
    return _form_view(request, SupplierForm, "Добавить поставщика", "Поставщик сохранён.", "finance:supplier_list")


@finance_staff_required
def supplier_edit(request, pk):
    return _form_view(request, SupplierForm, "Изменить поставщика", "Поставщик обновлён.", "finance:supplier_list", get_object_or_404(Supplier, pk=pk))


@finance_staff_required
def supplier_delete(request, pk):
    return _delete_view(request, get_object_or_404(Supplier, pk=pk), "Поставщик", "finance:supplier_list")


@finance_staff_required
def supplier_detail(request, pk):
    supplier = get_object_or_404(Supplier, pk=pk)
    start, end, context = _period_context(request)
    items = PartItem.objects.filter(receipt__supplier=supplier, receipt__date__range=(start, end)).select_related("receipt__part").prefetch_related("warranty_claims")
    context.update(supplier=supplier, stats=supplier_summary(supplier, start, end), page_obj=_page(items, request))
    return render(request, "finance/supplier_detail.html", context)


@finance_staff_required
def supplier_payment_list(request):
    start, end, context = _period_context(request)
    qs = SupplierPayment.objects.select_related("supplier", "receipt").filter(date__range=(start, end))
    if request.GET.get("supplier"):
        qs = qs.filter(supplier_id=request.GET["supplier"])
    context.update(page_obj=_page(qs, request))
    return render(request, "finance/supplier_payment_list.html", context)


@finance_staff_required
def supplier_payment_create(request):
    initial = {"supplier": request.GET.get("supplier")} if request.GET.get("supplier") else None
    return _form_view(request, SupplierPaymentForm, "Оплата поставщику", "Оплата сохранена.", "finance:supplier_payment_list", initial=initial)


@finance_staff_required
def supplier_payment_edit(request, pk):
    return _form_view(request, SupplierPaymentForm, "Изменить оплату поставщику", "Оплата обновлена.", "finance:supplier_payment_list", get_object_or_404(SupplierPayment, pk=pk))


@finance_staff_required
def supplier_return_list(request):
    supplier_id = request.GET.get("supplier")
    items = SupplierReturn.objects.select_related("supplier", "part_item__receipt__part").all()
    if supplier_id:
        items = items.filter(supplier_id=supplier_id)
    return render(request, "finance/supplier_return_list.html", {"items": items, "supplier_id": supplier_id})


@finance_staff_required
def supplier_return_create(request):
    initial = {"supplier": request.GET.get("supplier")} if request.GET.get("supplier") else None
    return _form_view(request, SupplierReturnForm, "Возврат поставщику", "Возврат сохранён.", "finance:supplier_return_list", initial=initial)


@finance_staff_required
def supplier_return_edit(request, pk):
    return _form_view(request, SupplierReturnForm, "Изменить возврат", "Возврат обновлён.", "finance:supplier_return_list", get_object_or_404(SupplierReturn, pk=pk))


def _reconciliation_context(request, supplier):
    preset = request.GET.get("period", "month")
    if preset == "all":
        start = end = None
    else:
        start, end, preset = resolve_period(request)
    query = (request.GET.get("q") or "").strip()
    operation_type = request.GET.get("type", "all")
    ledger = supplier_ledger(supplier, start, end, query, operation_type)
    return {
        "supplier": supplier, "rows": ledger["rows"], "summary": ledger["summary"],
        "date_from": start, "date_to": end, "period": preset, "q": query,
        "operation_type": operation_type,
        "operation_types": [("all", "Все операции"), ("receipt", "Поставки"), ("payment", "Оплаты"), ("return", "Возвраты"), ("warranty", "Гарантия"), ("replacement", "Замены"), ("adjustment", "Корректировки")],
    }


@finance_staff_required
def supplier_reconciliation(request, pk):
    supplier = get_object_or_404(Supplier, pk=pk)
    return render(request, "finance/supplier_reconciliation.html", _reconciliation_context(request, supplier))


@finance_staff_required
def supplier_reconciliation_excel(request, pk):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from django.utils.text import slugify

    supplier = get_object_or_404(Supplier, pk=pk)
    context = _reconciliation_context(request, supplier)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сверка"
    sheet.append(["ТЕХСФЕРА"])
    sheet.append(["Сверка с поставщиком"])
    sheet.append(["Поставщик", supplier.name])
    period_text = "Всё время" if not context["date_from"] else f'{context["date_from"]:%d.%m.%Y} — {context["date_to"]:%d.%m.%Y}'
    sheet.append(["Период", period_text])
    sheet.append(["Дата формирования", timezone.localdate()])
    sheet.append(["Начальное сальдо", context["summary"]["opening"]])
    sheet.append([])
    headers = ["Дата", "Тип операции", "Документ", "Деталь", "Модель", "Количество", "Цена", "Приход", "Возврат/зачёт", "Оплата", "Корректировка", "Долг после операции", "Комментарий"]
    sheet.append(headers)
    header_row = sheet.max_row
    for row in context["rows"]:
        sheet.append([row.date, row.operation_label, row.document, row.description, row.device_model, row.quantity, row.unit_price, row.debit or None, row.return_amount or None, row.payment or None, row.adjustment or None, row.balance, row.comment])
    summary = context["summary"]
    sheet.append([])
    sheet.append(["ИТОГО"])
    for label, value in (("Начальное сальдо", summary["opening"]), ("Поставки", summary["received"]), ("Оплаты", summary["paid"]), ("Возвраты/зачёты", summary["returns"]), ("Гарантийные корректировки", summary["warranty_adjustments"]), ("Другие корректировки", summary["adjustments"]), ("Конечное сальдо", summary["closing"])):
        sheet.append([label, value])
    sheet["A1"].font = Font(bold=True, size=16)
    sheet["A2"].font = Font(bold=True, size=14)
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:M{header_row + len(context['rows'])}"
    widths = [13, 23, 18, 30, 24, 12, 14, 14, 17, 14, 17, 20, 42]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=header_row + 1):
        row[0].number_format = "DD.MM.YYYY"
        row[12].alignment = Alignment(wrap_text=True, vertical="top")
        for index in range(6, 12):
            row[index].number_format = '#,##0.00 "BYN"'
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    start_name = context["date_from"].isoformat() if context["date_from"] else "all"
    end_name = context["date_to"].isoformat() if context["date_to"] else "all"
    response["Content-Disposition"] = f'attachment; filename="Supplier_Reconciliation_{slugify(supplier.name) or supplier.pk}_{start_name}_{end_name}.xlsx"'
    workbook.save(response)
    return response


@finance_staff_required
def supplier_payment_delete(request, pk):
    return _delete_view(request, get_object_or_404(SupplierPayment, pk=pk), "Оплата поставщику", "finance:supplier_payment_list")


@finance_staff_required
def warranty_list(request):
    query = (request.GET.get("q") or "").strip()
    qs = WarrantyClaim.objects.select_related("part_item__receipt__part", "part_item__receipt__supplier")
    if query:
        qs = qs.filter(Q(reason__icontains=query) | Q(defect_description__icontains=query) | Q(part_item__receipt__part__device_model__icontains=query) | Q(part_item__receipt__supplier__name__icontains=query))
    return render(request, "finance/warranty_list.html", {"page_obj": _page(qs, request), "q": query})


@finance_staff_required
def warranty_create(request):
    initial = {"part_item": request.GET.get("part_item")} if request.GET.get("part_item") else None
    return _form_view(request, WarrantyClaimForm, "Открыть гарантийный случай", "Гарантийный случай открыт.", "finance:warranty_list", initial=initial)


@finance_staff_required
def warranty_detail(request, pk):
    claim = get_object_or_404(
        WarrantyClaim.objects.select_related(
            "part_item__receipt__part", "part_item__receipt__supplier",
            "part_item__repair_usage__repair__employee",
        ), pk=pk,
    )
    return render(request, "finance/warranty_detail.html", {"claim": claim})


@finance_staff_required
def warranty_edit(request, pk):
    return _form_view(request, WarrantyClaimForm, "Изменить гарантийный случай", "Гарантийный случай обновлён.", "finance:warranty_list", get_object_or_404(WarrantyClaim, pk=pk))


@finance_staff_required
def distributed_expense_list(request):
    items = DistributedExpense.objects.select_related("category", "source_expense").prefetch_related("employees")
    return render(request, "finance/distributed_expense_list.html", {"items": items})


@finance_staff_required
def distributed_expense_create(request):
    return _form_view(request, DistributedExpenseForm, "Распределяемый расход", "Распределение сохранено.", "finance:distributed_expense_list")


@finance_staff_required
def distributed_expense_edit(request, pk):
    item = get_object_or_404(DistributedExpense, pk=pk)
    if item.allocations.exists():
        messages.info(request, "Изменение не затронет уже закрытые расчёты: они хранят snapshot.")
    return _form_view(request, DistributedExpenseForm, "Изменить распределение", "Распределение обновлено.", "finance:distributed_expense_list", item)


@finance_staff_required
def distributed_expense_detail(request, pk):
    item = get_object_or_404(DistributedExpense.objects.prefetch_related("employees", "allocations__calculation__period", "allocations__calculation__employee"), pk=pk)
    return render(request, "finance/distributed_expense_detail.html", {"item": item})


@finance_staff_required
def distributed_expense_finish(request, pk):
    item = get_object_or_404(DistributedExpense, pk=pk)
    if request.method == "POST":
        item.is_active = False
        item.finished_at = timezone.localdate()
        item.save(update_fields=["is_active", "finished_at"])
        messages.success(request, f"Распределение завершено. Нераспределённый остаток: {item.remaining_amount} BYN.")
    return redirect("finance:distributed_expense_detail", pk=pk)


@finance_staff_required
def payroll_period_list(request):
    return render(request, "finance/payroll_period_list.html", {"periods": PayrollPeriod.objects.prefetch_related("calculations__employee")})


@finance_staff_required
def payroll_period_create(request):
    return _form_view(request, PayrollPeriodForm, "Новый зарплатный период", "Период создан.", "finance:payroll_period_list")


@finance_staff_required
def payroll_period_detail(request, pk):
    period = get_object_or_404(PayrollPeriod, pk=pk)
    employee_id = request.GET.get("employee")
    if period.status == PayrollPeriod.Status.CLOSED:
        rows = period.calculations.select_related("employee").prefetch_related("allocations__distributed_expense", "repair_snapshots__repair")
    else:
        employees = Employee.objects.filter(is_active=True, is_owner=False)
        if employee_id:
            employees = employees.filter(pk=employee_id)
        rows = [{"employee": employee, **payroll_preview(period, employee)} for employee in employees]
    if employee_id and period.status == PayrollPeriod.Status.CLOSED:
        rows = rows.filter(employee_id=employee_id)
    return render(request, "finance/payroll_period_detail.html", {"period": period, "rows": rows})


@finance_staff_required
def payroll_period_close(request, pk):
    period = get_object_or_404(PayrollPeriod, pk=pk)
    if request.method == "POST":
        close_payroll_period(period)
        messages.success(request, "Расчётный период закрыт. Все суммы зафиксированы.")
    return redirect("finance:payroll_period_detail", pk=pk)
