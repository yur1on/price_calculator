from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db.models import Sum
from django.test import TestCase
from django.urls import reverse

from .models import (
    Employee, Expense, ExpenseCategory, OtherIncome, PartCatalog, PartItem,
    RepairFinance, RepairPart, SalaryPayment, StockReceipt, Supplier,
    SupplierPayment, SupplierReturn, WarrantyClaim,
    DistributedExpense, DistributedExpenseAllocation, PayrollCalculation, PayrollPeriod,
    PayrollRepairSnapshot,
)
from .services import close_payroll_period, employee_rows, financial_summary, payroll_preview, supplier_summary
from .supplier_ledger import supplier_ledger


DAY = date(2026, 9, 19)


class RepairCalculationTests(TestCase):
    def setUp(self):
        self.owner = Employee.objects.create(name="Юрий", default_percent="50.00", is_owner=True)
        self.master = Employee.objects.create(name="Сергей", default_percent="35.00")

    def repair(self, employee, revenue, parts, percent=None):
        return RepairFinance.objects.create(
            date=DAY, description="Тестовый ремонт", revenue=revenue, part_cost=parts,
            employee=employee,
            master_percent=employee.default_percent if percent is None else percent,
        )

    def test_owner_repair_has_no_salary(self):
        repair = self.repair(self.owner, "250.00", "100.00")
        self.assertEqual(repair.repair_margin, Decimal("150.00"))
        self.assertEqual(repair.master_salary, Decimal("0.00"))
        self.assertEqual(repair.workshop_profit, Decimal("150.00"))
        self.assertEqual(repair.master_percent, Decimal("0.00"))

    def test_master_repair_35_percent(self):
        repair = self.repair(self.master, "250.00", "100.00")
        self.assertEqual(repair.repair_margin, Decimal("150.00"))
        self.assertEqual(repair.master_salary, Decimal("52.50"))
        self.assertEqual(repair.workshop_profit, Decimal("97.50"))

    def test_repair_without_part(self):
        repair = self.repair(self.master, "100.00", "0.00")
        self.assertEqual(repair.master_salary, Decimal("35.00"))
        self.assertEqual(repair.workshop_profit, Decimal("65.00"))

    def test_loss_repair_salary_is_zero(self):
        repair = self.repair(self.master, "100.00", "150.00")
        self.assertEqual(repair.repair_margin, Decimal("-50.00"))
        self.assertEqual(repair.master_salary, Decimal("0.00"))
        self.assertEqual(repair.workshop_profit, Decimal("-50.00"))

    def test_employee_percent_change_does_not_change_old_repair(self):
        repair = self.repair(self.master, "250.00", "100.00")
        self.master.default_percent = Decimal("40.00")
        self.master.save()
        repair.refresh_from_db()
        self.assertEqual(repair.master_percent, Decimal("35.00"))
        self.assertEqual(repair.master_salary, Decimal("52.50"))

    def test_negative_money_and_invalid_percent_are_rejected(self):
        with self.assertRaises(ValidationError):
            self.repair(self.master, "-1.00", "0.00")
        with self.assertRaises(ValidationError):
            self.repair(self.master, "1.00", "0.00", "101.00")


class FinanceSummaryTests(TestCase):
    def setUp(self):
        self.master = Employee.objects.create(name="Сергей", default_percent="35.00")
        self.category, _ = ExpenseCategory.objects.get_or_create(name="Аренда")

    def test_net_profit_incomes_expenses_and_category(self):
        RepairFinance.objects.create(
            date=DAY, description="Ремонт", revenue="250", part_cost="100",
            employee=self.master, master_percent="35",
        )
        Expense.objects.create(date=DAY, amount="40", category=self.category, description="Аренда")
        OtherIncome.objects.create(date=DAY, amount="20", description="Продажа")
        summary = financial_summary(DAY, DAY)
        self.assertEqual(summary["revenue"], Decimal("250.00"))
        self.assertEqual(summary["parts"], Decimal("100.00"))
        self.assertEqual(summary["salary"], Decimal("52.50"))
        self.assertEqual(summary["expenses"], Decimal("40.00"))
        self.assertEqual(summary["other_income"], Decimal("20.00"))
        self.assertEqual(summary["net_profit"], Decimal("77.50"))
        self.assertEqual(
            Expense.objects.filter(category=self.category).aggregate(v=Sum("amount"))["v"],
            Decimal("40.00"),
        )

    def test_salary_debt_and_payment_not_in_net_profit(self):
        RepairFinance.objects.create(
            date=DAY, description="Ремонт", revenue="250", part_cost="100",
            employee=self.master, master_percent="35",
        )
        before = financial_summary(DAY, DAY)["net_profit"]
        SalaryPayment.objects.create(employee=self.master, date=DAY, amount="20")
        row = employee_rows(DAY, DAY, self.master)[0]
        self.assertEqual(row["salary"], Decimal("52.50"))
        self.assertEqual(row["paid"], Decimal("20.00"))
        self.assertEqual(row["debt"], Decimal("32.50"))
        self.assertEqual(financial_summary(DAY, DAY)["net_profit"], before)


class FinanceAccessTests(TestCase):
    def test_finance_requires_staff(self):
        user = get_user_model().objects.create_user(username="client", password="test")
        self.client.force_login(user)
        response = self.client.get(reverse("finance:dashboard"))
        self.assertEqual(response.status_code, 302)

    def test_staff_can_open_dashboard(self):
        user = get_user_model().objects.create_superuser(username="owner", password="test")
        self.client.force_login(user)
        response = self.client.get(reverse("finance:dashboard"))
        self.assertEqual(response.status_code, 200)

    def test_finance_menu_is_hidden_from_public_and_visible_to_admin(self):
        public_response = self.client.get(reverse("home"))
        self.assertNotContains(public_response, "💰 Финансы")
        admin = get_user_model().objects.create_superuser(username="admin", password="test")
        self.client.force_login(admin)
        admin_response = self.client.get(reverse("home"))
        self.assertContains(admin_response, "💰 Финансы")
        self.assertContains(admin_response, reverse("finance:dashboard"))


class FinanceCrudTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(username="admin", password="test")
        self.client.force_login(self.admin)
        self.master = Employee.objects.create(name="Сергей", default_percent="35.00")
        self.category, _ = ExpenseCategory.objects.get_or_create(name="Расходники")

    def test_all_finance_sections_open(self):
        names = [
            "dashboard", "repair_list", "expense_list", "income_list", "master_list",
            "salary_list", "category_list", "stock_list", "receipt_list", "catalog_list",
            "supplier_list", "supplier_payment_list", "warranty_list",
        ]
        for name in names:
            with self.subTest(name=name):
                self.assertEqual(self.client.get(reverse(f"finance:{name}")).status_code, 200)

    def test_new_repair_form_contains_today_in_html_date_format(self):
        response = self.client.get(reverse("finance:repair_create"))
        self.assertContains(response, f'value="{date.today():%Y-%m-%d}"')

    def test_repair_crud(self):
        response = self.client.post(reverse("finance:repair_create"), {
            "date": DAY, "description": "iPhone 11 — дисплей", "revenue": "250.00",
            "manual_part_cost": "100.00", "employee": self.master.pk, "master_percent": "35.00", "comment": "",
        })
        self.assertRedirects(response, reverse("finance:repair_list"))
        repair = RepairFinance.objects.get()
        self.assertEqual(repair.workshop_profit, Decimal("97.50"))
        self.client.post(reverse("finance:repair_edit", args=[repair.pk]), {
            "date": DAY, "description": "iPhone 11 — дисплей OLED", "revenue": "260.00",
            "manual_part_cost": "100.00", "employee": self.master.pk, "master_percent": "35.00", "comment": "",
        })
        repair.refresh_from_db()
        self.assertIn("OLED", repair.description)
        self.client.post(reverse("finance:repair_delete", args=[repair.pk]))
        self.assertFalse(RepairFinance.objects.exists())

    def test_expense_income_master_and_salary_crud(self):
        self.client.post(reverse("finance:expense_create"), {
            "date": DAY, "category": self.category.pk, "description": "Флюс", "amount": "45.00", "comment": "",
        })
        self.client.post(reverse("finance:income_create"), {
            "date": DAY, "description": "Продажа запчасти", "amount": "30.00", "comment": "",
        })
        self.client.post(reverse("finance:salary_payment_create"), {
            "date": DAY, "employee": self.master.pk, "amount": "20.00", "comment": "Аванс",
        })
        self.assertEqual(Expense.objects.count(), 1)
        self.assertEqual(OtherIncome.objects.count(), 1)
        self.assertEqual(SalaryPayment.objects.count(), 1)
        response = self.client.post(reverse("finance:master_create"), {
            "name": "Иван", "default_percent": "40.00", "is_active": "on",
        })
        self.assertRedirects(response, reverse("finance:master_list"))
        self.assertTrue(Employee.objects.filter(name="Иван", default_percent="40.00").exists())


class InventoryTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(username="stock-admin", password="test")
        self.client.force_login(self.admin)
        self.supplier = Supplier.objects.create(name="MobileParts", contact_person="Александр")
        self.part = PartCatalog.objects.create(brand="Samsung", device_model="A55", name="OLED дисплей")
        self.master = Employee.objects.create(name="Сергей", default_percent="35.00")

    def receipt(self, quantity=3, cost="180.00"):
        return StockReceipt.objects.create(
            date=DAY, supplier=self.supplier, part=self.part, quantity=quantity,
            unit_cost=cost, order_number="MP-4582",
        )

    def test_supplier_and_receipt_create_concrete_stock_items(self):
        receipt = self.receipt()
        self.assertEqual(receipt.items.count(), 3)
        self.assertEqual(set(receipt.items.values_list("status", flat=True)), {PartItem.Status.IN_STOCK})
        self.assertEqual(receipt.items.first().inventory_code[:1], "P")

    def test_receipt_form_accepts_manually_typed_part(self):
        response = self.client.post(reverse("finance:receipt_create"), {
            "date": DAY, "supplier": self.supplier.pk,
            "part_name": "Samsung — A60 — дисплей", "quantity": 2,
            "unit_cost": "75.00", "order_number": "MANUAL-1", "comment": "",
        })
        self.assertRedirects(response, reverse("finance:receipt_list"))
        receipt = StockReceipt.objects.get(order_number="MANUAL-1")
        self.assertEqual(receipt.part.brand, "Samsung")
        self.assertEqual(receipt.part.device_model, "A60")
        self.assertEqual(receipt.part.name, "дисплей")
        self.assertEqual(receipt.items.count(), 2)

    def test_stock_parts_are_installed_and_summed_in_repair(self):
        items = list(self.receipt(quantity=2, cost="100.00").items.all())
        response = self.client.post(reverse("finance:repair_create"), {
            "date": DAY, "description": "Samsung A55", "revenue": "300.00",
            "manual_part_cost": "5.00", "part_items": [item.pk for item in items],
            "employee": self.master.pk, "master_percent": "35.00", "comment": "",
        })
        self.assertRedirects(response, reverse("finance:repair_list"))
        repair = RepairFinance.objects.get()
        self.assertEqual(repair.part_cost, Decimal("205.00"))
        self.assertEqual(repair.master_salary, Decimal("33.25"))
        self.assertEqual(repair.workshop_profit, Decimal("61.75"))
        self.assertEqual(set(PartItem.objects.values_list("status", flat=True)), {PartItem.Status.INSTALLED})
        self.assertEqual(repair.repair_parts.count(), 2)

    def test_stock_part_repair_saves_with_empty_manual_cost(self):
        item = self.receipt(quantity=1, cost="120.00").items.get()
        response = self.client.post(reverse("finance:repair_create"), {
            "date": DAY, "description": "Samsung A40", "revenue": "140.00",
            "manual_part_cost": "", "part_items": [item.pk],
            "employee": self.master.pk, "master_percent": "35.00", "comment": "",
        })
        self.assertRedirects(response, reverse("finance:repair_list"))
        repair = RepairFinance.objects.get()
        self.assertEqual(repair.manual_part_cost, Decimal("0.00"))
        self.assertEqual(repair.part_cost, Decimal("120.00"))
        self.assertEqual(repair.master_salary, Decimal("7.00"))
        self.assertEqual(repair.workshop_profit, Decimal("13.00"))

    def test_part_item_cannot_be_used_twice(self):
        item = self.receipt(quantity=1).items.get()
        first = RepairFinance.objects.create(date=DAY, description="Первый", revenue="300", part_cost="0", employee=self.master, master_percent="35")
        second = RepairFinance.objects.create(date=DAY, description="Второй", revenue="300", part_cost="0", employee=self.master, master_percent="35")
        RepairPart.objects.create(repair=first, part_item=item, cost_used=item.unit_cost, installed_at=DAY)
        item.refresh_from_db()
        with self.assertRaises(ValidationError):
            RepairPart.objects.create(repair=second, part_item=item, cost_used=item.unit_cost, installed_at=DAY)

    def test_old_manual_repair_still_works(self):
        repair = RepairFinance.objects.create(date=DAY, description="Старый ремонт", revenue="250", part_cost="100", employee=self.master, master_percent="35")
        self.assertEqual(repair.manual_part_cost, Decimal("100.00"))
        self.assertEqual(repair.workshop_profit, Decimal("97.50"))

    def test_warranty_is_linked_to_physical_item(self):
        item = self.receipt(quantity=1).items.get()
        repair = RepairFinance.objects.create(date=DAY, description="Ремонт", revenue="300", part_cost="0", employee=self.master, master_percent="35")
        RepairPart.objects.create(repair=repair, part_item=item, cost_used=item.unit_cost, installed_at=DAY)
        claim = WarrantyClaim.objects.create(part_item=item, opened_at=DAY, reason="Полосы", defect_description="Полосы на экране")
        item.refresh_from_db()
        self.assertEqual(claim.part_item, item)
        self.assertEqual(item.status, PartItem.Status.WARRANTY)

    def test_supplier_report_and_debt(self):
        self.receipt(quantity=3, cost="100.00")
        SupplierPayment.objects.create(supplier=self.supplier, date=DAY, amount="120.00")
        stats = supplier_summary(self.supplier, DAY, DAY)
        self.assertEqual(stats["received_count"], 3)
        self.assertEqual(stats["received_value"], Decimal("300.00"))
        self.assertEqual(stats["paid"], Decimal("120.00"))
        self.assertEqual(stats["debt"], Decimal("180.00"))

    def test_supplier_overpayment_becomes_positive_balance(self):
        self.receipt(quantity=1, cost="100.00")
        SupplierPayment.objects.create(supplier=self.supplier, date=DAY, amount="140.00")
        stats = supplier_summary(self.supplier, DAY, DAY)
        self.assertEqual(stats["debt"], Decimal("0.00"))
        self.assertEqual(stats["credit"], Decimal("40.00"))
        self.assertEqual(stats["balance"], Decimal("40.00"))

    def test_supplier_overview_clearly_separates_debt_and_advance(self):
        self.receipt(quantity=1, cost="100.00")
        SupplierPayment.objects.create(supplier=self.supplier, date=DAY, amount="40.00")
        prepaid = Supplier.objects.create(name="PartsCity")
        SupplierPayment.objects.create(supplier=prepaid, date=DAY, amount="25.00")
        response = self.client.get(reverse("finance:supplier_list"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_debt"], Decimal("60.00"))
        self.assertEqual(response.context["total_credit"], Decimal("25.00"))
        self.assertContains(response, "Вы должны поставщику")
        self.assertContains(response, "Поставщик должен вам / ваш аванс")

    def test_repair_with_installed_part_cannot_be_deleted(self):
        item = self.receipt(quantity=1).items.get()
        repair = RepairFinance.objects.create(date=DAY, description="Ремонт", revenue="300", part_cost="0", employee=self.master, master_percent="35")
        RepairPart.objects.create(repair=repair, part_item=item, cost_used=item.unit_cost, installed_at=DAY)
        response = self.client.post(reverse("finance:repair_delete", args=[repair.pk]))
        self.assertRedirects(response, reverse("finance:repair_list"))
        self.assertTrue(RepairFinance.objects.filter(pk=repair.pk).exists())
        item.refresh_from_db()
        self.assertEqual(item.status, PartItem.Status.INSTALLED)

    def test_inventory_access_is_protected(self):
        self.client.logout()
        user = get_user_model().objects.create_user(username="visitor", password="test")
        self.client.force_login(user)
        for name in ("stock_list", "receipt_list", "supplier_list", "warranty_list"):
            self.assertEqual(self.client.get(reverse(f"finance:{name}")).status_code, 302)


class DistributedPayrollTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(username="payroll-admin", password="test")
        self.client.force_login(self.admin)
        self.master = Employee.objects.create(name="Сергей", default_percent="35.00")
        self.other = Employee.objects.create(name="Иван", default_percent="40.00")
        self.category, _ = ExpenseCategory.objects.get_or_create(name="Расходные материалы")
        self.period = PayrollPeriod.objects.create(start_date=date(2026, 9, 1), end_date=date(2026, 9, 14))
        RepairFinance.objects.create(
            date=date(2026, 9, 5), description="Ремонты периода", revenue="3000.00",
            part_cost="800.00", employee=self.master, master_percent="35.00",
        )
        self.real_expense = Expense.objects.create(
            date=date(2026, 9, 1), amount="180.00", category=self.category, description="BGA-флюс",
        )
        self.distributed = DistributedExpense.objects.create(
            name="BGA-флюс", category=self.category, source_expense=self.real_expense,
            start_date=date(2026, 9, 1), total_cost="180.00", periods_count=6,
        )
        self.distributed.employees.add(self.master)

    def test_per_period_amount_and_payroll_formula(self):
        self.assertEqual(self.distributed.per_period_amount, Decimal("30.00"))
        data = payroll_preview(self.period, self.master)
        self.assertEqual(data["repair_margin"], Decimal("2200.00"))
        self.assertEqual(data["distributed_costs"], Decimal("30.00"))
        self.assertEqual(data["salary_base"], Decimal("2170.00"))
        self.assertEqual(data["salary_amount"], Decimal("759.50"))

    def test_multiple_distributions_are_summed_and_employee_specific(self):
        second = DistributedExpense.objects.create(
            name="Оплётка", category=self.category, start_date=date(2026, 9, 1),
            total_cost="60.00", periods_count=6,
        )
        second.employees.add(self.master)
        self.assertEqual(payroll_preview(self.period, self.master)["distributed_costs"], Decimal("40.00"))
        self.assertEqual(payroll_preview(self.period, self.other)["distributed_costs"], Decimal("0.00"))

    def test_distributed_expense_does_not_reduce_net_profit_twice(self):
        summary = financial_summary(date(2026, 9, 1), date(2026, 9, 14))
        expected = Decimal("2200.00") - Decimal("770.00") - Decimal("180.00")
        self.assertEqual(summary["net_profit"], expected)

    def test_close_creates_immutable_snapshot_and_cannot_allocate_twice(self):
        close_payroll_period(self.period)
        calc = PayrollCalculation.objects.get(period=self.period, employee=self.master)
        self.assertEqual(calc.salary_amount, Decimal("759.50"))
        self.assertEqual(DistributedExpenseAllocation.objects.filter(calculation=calc).count(), 1)
        self.distributed.total_cost = Decimal("600.00")
        self.distributed.save()
        calc.refresh_from_db()
        self.assertEqual(calc.salary_amount, Decimal("759.50"))
        close_payroll_period(self.period)
        self.assertEqual(PayrollCalculation.objects.filter(period=self.period, employee=self.master).count(), 1)

    def test_multiple_masters_split_one_period_share(self):
        self.distributed.employees.add(self.other)
        close_payroll_period(self.period)
        allocations = self.distributed.allocations.filter(calculation__period=self.period)
        self.assertEqual(allocations.count(), 2)
        self.assertEqual(sum((row.amount for row in allocations), Decimal("0.00")), Decimal("30.00"))

    def test_payroll_pages_are_protected(self):
        client = self.client_class()
        user = get_user_model().objects.create_user(username="no-finance", password="test")
        client.force_login(user)
        self.assertEqual(client.get(reverse("finance:payroll_period_list")).status_code, 302)
        self.assertEqual(client.get(reverse("finance:distributed_expense_list")).status_code, 302)

    def test_preview_lists_only_selected_master_repairs_inside_period(self):
        own = RepairFinance.objects.create(date=date(2026, 9, 7), description="Ремонт Сергея", revenue="200", part_cost="50", employee=self.master, master_percent="35")
        RepairFinance.objects.create(date=date(2026, 9, 7), description="Ремонт Ивана", revenue="900", part_cost="10", employee=self.other, master_percent="40")
        RepairFinance.objects.create(date=date(2026, 9, 20), description="Ремонт вне периода", revenue="700", part_cost="20", employee=self.master, master_percent="35")
        data = payroll_preview(self.period, self.master)
        self.assertEqual([row["repair"] for row in data["repairs"]], [RepairFinance.objects.get(description="Ремонты периода"), own])
        self.assertEqual(data["repairs_count"], 2)
        self.assertEqual(data["revenue"], Decimal("3200.00"))
        self.assertEqual(data["direct_costs"], Decimal("850.00"))
        response = self.client.get(reverse("finance:payroll_period_detail", args=[self.period.pk]), {"employee": self.master.pk})
        self.assertContains(response, "Ремонт Сергея")
        self.assertNotContains(response, "Ремонт Ивана")
        self.assertNotContains(response, "Ремонт вне периода")
        self.assertContains(response, "Распределяемые расходы периода")

    def test_closed_period_keeps_historical_repair_rows(self):
        repair = RepairFinance.objects.get(description="Ремонты периода")
        close_payroll_period(self.period)
        snapshot = PayrollRepairSnapshot.objects.get(repair=repair)
        self.assertEqual(snapshot.revenue, Decimal("3000.00"))
        self.assertEqual(snapshot.parts_cost, Decimal("800.00"))
        repair.description = "Изменённый ремонт"
        repair.revenue = Decimal("10.00")
        repair.part_cost = Decimal("0.00")
        repair.save()
        response = self.client.get(reverse("finance:payroll_period_detail", args=[self.period.pk]), {"employee": self.master.pk})
        self.assertContains(response, "Ремонты периода")
        self.assertNotContains(response, "Изменённый ремонт")
        snapshot.refresh_from_db()
        self.assertEqual(snapshot.revenue, Decimal("3000.00"))
        repair.delete()
        snapshot.refresh_from_db()
        self.assertIsNone(snapshot.repair)
        self.assertEqual(snapshot.description, "Ремонты периода")

    def test_distribution_progress_changes_only_when_period_closes(self):
        first_preview = payroll_preview(self.period, self.master)["distributions"][0]
        self.assertEqual(first_preview["part_number"], 1)
        self.assertEqual(first_preview["accounted_after"], Decimal("30.00"))
        self.assertEqual(first_preview["remaining_after"], Decimal("150.00"))
        payroll_preview(self.period, self.master)
        payroll_preview(self.period, self.master)
        self.assertEqual(self.distributed.used_periods, 0)

        close_payroll_period(self.period)
        first = self.distributed.allocations.get(calculation__period=self.period)
        self.assertEqual((first.part_number, first.parts_count), (1, 6))
        self.assertEqual(first.accounted_after, Decimal("30.00"))
        self.assertEqual(first.remaining_after, Decimal("150.00"))

        second_period = PayrollPeriod.objects.create(start_date=date(2026, 9, 15), end_date=date(2026, 9, 28))
        second_preview = payroll_preview(second_period, self.master)["distributions"][0]
        self.assertEqual(second_preview["part_number"], 2)
        self.assertEqual(self.distributed.used_periods, 1)
        close_payroll_period(second_period)
        second = self.distributed.allocations.get(calculation__period=second_period)
        self.assertEqual(second.part_number, 2)
        self.assertEqual(second.accounted_after, Decimal("60.00"))
        first.refresh_from_db()
        self.assertEqual(first.part_number, 1)

    def test_distribution_finishes_after_six_parts(self):
        close_payroll_period(self.period)
        for number in range(2, 7):
            start = date(2026, 9, 1) + timedelta(days=(number - 1) * 14)
            period = PayrollPeriod.objects.create(start_date=start, end_date=start + timedelta(days=13))
            close_payroll_period(period)
        self.distributed.refresh_from_db()
        self.assertFalse(self.distributed.is_active)
        allocations = self.distributed.allocations.order_by("part_number")
        self.assertEqual(list(allocations.values_list("part_number", flat=True)), [1, 2, 3, 4, 5, 6])
        self.assertEqual(allocations.last().accounted_after, Decimal("180.00"))
        self.assertEqual(allocations.last().remaining_after, Decimal("0.00"))
        next_period = PayrollPeriod.objects.create(start_date=date(2026, 12, 1), end_date=date(2026, 12, 14))
        self.assertEqual(payroll_preview(next_period, self.master)["distributions"], [])

    def test_rounding_remainder_and_independent_progress(self):
        rounded = DistributedExpense.objects.create(
            name="Паста", category=self.category, start_date=date(2026, 9, 1),
            total_cost="100.00", periods_count=3,
        )
        rounded.employees.add(self.master)
        close_payroll_period(self.period)
        for offset in (14, 28):
            start = date(2026, 9, 1) + timedelta(days=offset)
            close_payroll_period(PayrollPeriod.objects.create(start_date=start, end_date=start + timedelta(days=13)))
        amounts = list(rounded.allocations.order_by("part_number").values_list("amount", flat=True))
        self.assertEqual(amounts, [Decimal("33.33"), Decimal("33.33"), Decimal("33.34")])
        self.assertEqual(sum(amounts, Decimal("0.00")), Decimal("100.00"))
        self.assertEqual(rounded.allocations.last().remaining_after, Decimal("0.00"))
        self.assertEqual(self.distributed.used_periods, 3)
        self.assertEqual(rounded.used_periods, 3)


class MasterPortalTests(TestCase):
    password = "SafeMaster-4826!"

    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(username="portal-admin", password="Admin-4826!")
        self.master = Employee.objects.create(name="Сергей", default_percent="35.00")
        self.other = Employee.objects.create(name="Иван", default_percent="40.00")
        self.user = get_user_model().objects.create_user(username="sergey", password=self.password)
        self.master.user = self.user
        self.master.save(update_fields=["user"])
        self.period = PayrollPeriod.objects.create(start_date=date(2026, 9, 1), end_date=date(2026, 9, 14))
        self.own_repair = RepairFinance.objects.create(date=date(2026, 9, 5), description="iPhone Сергея", revenue="300", part_cost="100", employee=self.master, master_percent="35")
        self.other_repair = RepairFinance.objects.create(date=date(2026, 9, 6), description="Samsung Ивана", revenue="500", part_cost="50", employee=self.other, master_percent="40")
        category, _ = ExpenseCategory.objects.get_or_create(name="Материалы кабинета")
        self.distributed = DistributedExpense.objects.create(name="BGA-флюс", category=category, start_date=date(2026, 9, 1), total_cost="180", periods_count=6)
        self.distributed.employees.add(self.master)

    def test_admin_creates_safe_master_account(self):
        employee = Employee.objects.create(name="Алексей", default_percent="35")
        self.client.force_login(self.admin)
        response = self.client.post(reverse("finance:master_access_create", args=[employee.pk]), {
            "username": "alexey", "password1": self.password, "password2": self.password,
        })
        self.assertRedirects(response, reverse("finance:master_list"))
        employee.refresh_from_db()
        self.assertFalse(employee.user.is_staff)
        self.assertFalse(employee.user.is_superuser)
        self.assertNotEqual(employee.user.password, self.password)
        self.assertTrue(employee.user.check_password(self.password))
        response = self.client.post(reverse("finance:master_access_create", args=[employee.pk]), {
            "username": "second-account", "password1": self.password, "password2": self.password,
        })
        self.assertRedirects(response, reverse("finance:master_list"))
        self.assertFalse(get_user_model().objects.filter(username="second-account").exists())

    def test_duplicate_username_is_rejected(self):
        employee = Employee.objects.create(name="Алексей", default_percent="35")
        self.client.force_login(self.admin)
        response = self.client.post(reverse("finance:master_access_create", args=[employee.pk]), {
            "username": "sergey", "password1": self.password, "password2": self.password,
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Пользователь с таким именем уже существует")
        employee.refresh_from_db()
        self.assertIsNone(employee.user)

    def test_login_dashboard_contains_only_own_data_and_distribution(self):
        response = self.client.post(reverse("master_portal:login"), {"username": "sergey", "password": self.password})
        self.assertRedirects(response, reverse("master_portal:dashboard"))
        response = self.client.get(reverse("master_portal:dashboard"))
        self.assertContains(response, "iPhone Сергея")
        self.assertNotContains(response, "Samsung Ивана")
        self.assertContains(response, "BGA-флюс")
        self.assertContains(response, "Будет:")
        self.assertContains(response, "1 / 6")
        self.assertEqual(self.distributed.used_periods, 0)

    def test_anonymous_and_unlinked_user_cannot_open_portal(self):
        self.assertRedirects(self.client.get(reverse("master_portal:dashboard")), f"{reverse('master_portal:login')}?next={reverse('master_portal:dashboard')}")
        client = self.client_class()
        client.force_login(get_user_model().objects.create_user(username="client", password="Client-4826!"))
        self.assertEqual(client.get(reverse("master_portal:dashboard")).status_code, 403)

    def test_master_cannot_access_or_mutate_finance(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("finance:dashboard")).status_code, 302)
        before = RepairFinance.objects.count()
        response = self.client.post(reverse("finance:repair_create"), {
            "date": DAY, "description": "Взлом", "revenue": "9999", "employee": self.master.pk,
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(RepairFinance.objects.count(), before)
        self.assertEqual(self.client.post(reverse("finance:payroll_period_close", args=[self.period.pk])).status_code, 302)
        self.assertEqual(self.period.status, PayrollPeriod.Status.OPEN)
        self.assertEqual(self.client.post(reverse("finance:salary_payment_create"), {"employee": self.master.pk, "amount": "999"}).status_code, 302)
        self.assertFalse(SalaryPayment.objects.exists())

    def test_closed_snapshot_history_and_idor_protection(self):
        close_payroll_period(self.period)
        SalaryPayment.objects.create(employee=self.master, date=date(2026, 9, 10), amount="40", comment="Аванс")
        self.own_repair.description = "Изменено позже"
        self.own_repair.revenue = Decimal("1")
        self.own_repair.save()
        self.client.force_login(self.user)
        response = self.client.get(reverse("master_portal:period_detail", args=[self.period.pk]))
        self.assertContains(response, "iPhone Сергея")
        self.assertNotContains(response, "Изменено позже")
        self.assertNotContains(response, "Samsung Ивана")
        self.assertContains(response, "1 / 6")
        self.assertContains(response, "Аванс")
        foreign_period = PayrollPeriod.objects.create(start_date=date(2026, 10, 1), end_date=date(2026, 10, 14), status=PayrollPeriod.Status.CLOSED)
        PayrollCalculation.objects.create(period=foreign_period, employee=self.other, repairs_count=0, revenue=0, direct_costs=0, repair_margin=0, distributed_costs=0, salary_base=0, percent=40, salary_amount=0)
        self.assertEqual(self.client.get(reverse("master_portal:period_detail", args=[foreign_period.pk])).status_code, 404)

    def test_access_toggle_and_password_change(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("finance:master_access_toggle", args=[self.master.pk]))
        self.user.refresh_from_db()
        self.assertFalse(self.user.is_active)
        self.client.logout()
        self.assertFalse(self.client.login(username="sergey", password=self.password))
        self.user.is_active = True
        self.user.save(update_fields=["is_active"])
        self.client.force_login(self.admin)
        new_password = "NewMaster-9357!"
        self.client.post(reverse("finance:master_access_password", args=[self.master.pk]), {"password1": new_password, "password2": new_password})
        self.user.refresh_from_db()
        self.assertFalse(self.user.check_password(self.password))
        self.assertTrue(self.user.check_password(new_password))

    def test_menu_visibility_and_no_public_registration(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("repairs:contacts"))
        self.assertContains(response, "💰 Моя зарплата")
        self.assertNotContains(response, "💰 Финансы")
        self.client.logout()
        response = self.client.get(reverse("repairs:contacts"))
        self.assertNotContains(response, "Моя зарплата")
        self.assertEqual(self.client.get("/register/").status_code, 404)


class SupplierReconciliationTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(username="recon-admin", password="test")
        self.client.force_login(self.admin)
        self.supplier = Supplier.objects.create(name="PartsMarket")
        self.iphone = PartCatalog.objects.create(brand="Apple", device_model="iPhone 11", name="OLED")
        self.samsung = PartCatalog.objects.create(brand="Samsung", device_model="A55", name="OLED")
        self.first = StockReceipt.objects.create(date=date(2026, 9, 1), supplier=self.supplier, part=self.iphone, quantity=2, unit_cost="120", order_number="НК-125")
        self.second = StockReceipt.objects.create(date=date(2026, 9, 1), supplier=self.supplier, part=self.samsung, quantity=1, unit_cost="150", order_number="НК-125")

    def test_invoice_items_are_separate_and_purchases_increase_debt(self):
        ledger = supplier_ledger(self.supplier)
        receipts = [row for row in ledger["rows"] if row.operation_type == "receipt"]
        self.assertEqual(len(receipts), 2)
        self.assertEqual([row.document for row in receipts], ["НК-125", "НК-125"])
        self.assertEqual([row.balance for row in receipts], [Decimal("240.00"), Decimal("390.00")])
        self.assertEqual(ledger["summary"]["current"], Decimal("390.00"))

    def test_partial_and_second_payment_reduce_balance(self):
        SupplierPayment.objects.create(supplier=self.supplier, date=date(2026, 9, 3), amount="100", payment_method="transfer")
        SupplierPayment.objects.create(supplier=self.supplier, date=date(2026, 9, 5), amount="140", payment_method="cash")
        ledger = supplier_ledger(self.supplier)
        payments = [row for row in ledger["rows"] if row.operation_type == "payment"]
        self.assertEqual([row.balance for row in payments], [Decimal("290.00"), Decimal("150.00")])
        self.assertEqual(ledger["summary"]["paid"], Decimal("240.00"))

    def test_physical_return_does_not_reduce_debt_until_confirmed(self):
        item = self.first.items.first()
        returned = SupplierReturn.objects.create(supplier=self.supplier, part_item=item, date=date(2026, 9, 7), reason="Брак", status=SupplierReturn.Status.SENT)
        self.assertEqual(supplier_ledger(self.supplier)["summary"]["current"], Decimal("390.00"))
        returned.status = SupplierReturn.Status.CREDITED
        returned.financial_date = date(2026, 9, 9)
        returned.financial_amount = Decimal("120.00")
        returned.save()
        ledger = supplier_ledger(self.supplier)
        self.assertEqual(ledger["summary"]["returns"], Decimal("120.00"))
        self.assertEqual(ledger["summary"]["current"], Decimal("270.00"))

    def test_warranty_and_replacement_have_no_fake_money(self):
        item = self.second.items.first()
        claim = WarrantyClaim.objects.create(part_item=item, opened_at=date(2026, 9, 4), reason="Полосы", defect_description="Полосы", sent_to_supplier_at=date(2026, 9, 6))
        before = supplier_ledger(self.supplier)["summary"]["current"]
        claim.status = WarrantyClaim.Status.REPLACED
        claim.closed_at = date(2026, 9, 10)
        claim.replacement_part_item = self.first.items.last()
        claim.save()
        ledger = supplier_ledger(self.supplier)
        self.assertEqual(ledger["summary"]["current"], before)
        self.assertTrue(any(row.operation_type == "replacement" for row in ledger["rows"]))

    def test_opening_closing_search_type_and_stable_order(self):
        SupplierPayment.objects.create(supplier=self.supplier, date=date(2026, 9, 1), amount="100", document_number="TX-1")
        third = StockReceipt.objects.create(date=date(2026, 9, 10), supplier=self.supplier, part=self.samsung, quantity=1, unit_cost="210", order_number="НК-132")
        period = supplier_ledger(self.supplier, date(2026, 9, 10), date(2026, 9, 30))
        self.assertEqual(period["summary"]["opening"], Decimal("290.00"))
        self.assertEqual(period["summary"]["closing"], Decimal("500.00"))
        self.assertEqual(period["rows"][0].description, third.part.name)
        search = supplier_ledger(self.supplier, query="A55")
        self.assertTrue(search["rows"])
        self.assertTrue(all("A55".lower() in row.search_text for row in search["rows"]))
        invoice = supplier_ledger(self.supplier, query="НК-125")
        self.assertEqual(len([row for row in invoice["rows"] if row.operation_type == "receipt"]), 2)
        only_payments = supplier_ledger(self.supplier, operation_type="payment")
        self.assertEqual({row.operation_type for row in only_payments["rows"]}, {"payment"})
        same_day = supplier_ledger(self.supplier, date(2026, 9, 1), date(2026, 9, 1))["rows"]
        self.assertEqual([row.operation_type for row in same_day], ["receipt", "receipt", "payment"])

    def test_installing_part_does_not_change_supplier_debt(self):
        master = Employee.objects.create(name="Мастер сверки", default_percent="35")
        repair = RepairFinance.objects.create(date=date(2026, 9, 8), description="Установка", revenue="300", part_cost="0", employee=master, master_percent="35")
        before = supplier_ledger(self.supplier)["summary"]["current"]
        RepairPart.objects.create(repair=repair, part_item=self.first.items.first(), cost_used="120", installed_at=date(2026, 9, 8))
        self.assertEqual(supplier_ledger(self.supplier)["summary"]["current"], before)

    def test_web_and_excel_match_selected_supplier_and_period(self):
        SupplierPayment.objects.create(supplier=self.supplier, date=date(2026, 9, 3), amount="100", document_number="PAY-1")
        params = {"period": "custom", "from": "2026-09-01", "to": "2026-09-30"}
        web = self.client.get(reverse("finance:supplier_reconciliation", args=[self.supplier.pk]), params)
        self.assertEqual(web.status_code, 200)
        self.assertContains(web, "PartsMarket")
        self.assertEqual([row.document for row in web.context["rows"]].count("НК-125"), 2)
        excel = self.client.get(reverse("finance:supplier_reconciliation_excel", args=[self.supplier.pk]), params)
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(excel["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        from openpyxl import load_workbook
        sheet = load_workbook(BytesIO(excel.content), data_only=True).active
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn("PartsMarket", values)
        self.assertIn("01.09.2026 — 30.09.2026", values)
        self.assertEqual(values.count("НК-125"), 2)
        self.assertIn("PAY-1", values)
        self.assertIn(Decimal("290.00"), [Decimal(str(value)) for value in values if isinstance(value, (int, float))])

    def test_reconciliation_is_protected_from_master_and_regular_user(self):
        master_user = get_user_model().objects.create_user(username="recon-master", password="test")
        employee = Employee.objects.create(name="Закрытый мастер", default_percent="35", user=master_user)
        url = reverse("finance:supplier_reconciliation", args=[self.supplier.pk])
        self.client.force_login(master_user)
        self.assertEqual(self.client.get(url).status_code, 302)
        self.client.force_login(get_user_model().objects.create_user(username="recon-client", password="test"))
        self.assertEqual(self.client.get(url).status_code, 302)
